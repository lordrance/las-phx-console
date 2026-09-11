# -*- coding: utf-8 -*-
"""
Build the LAS/PHX dashboard's data layer from the finished station table.

  mapdata.json            stations + dsp_summary, same contract as 北加 index.html
  price_adjustments.csv   zip,price -- the editable price sheet

DSP identity = fleet name with the station prefix stripped ('PHX-ECS' -> 'ECS'),
so one DSP operating at several stations shows up as one DSP. The original
fleet string is kept per zip for the popup.
"""
import csv
import json
import os
import re
import warnings
from collections import Counter, defaultdict

import openpyxl

warnings.filterwarnings("ignore")

BASE = r"C:\Users\uyran\OneDrive\Desktop\rance\nc-ops-console\模板和数据"
GEN = os.path.join(BASE, "拉斯维加斯凤凰城各站点信息.xlsx")
SRC = os.path.join(BASE, "GOFO WE DSP对应关系 （高链）.xlsx")
HERE = os.path.dirname(os.path.abspath(__file__))

# source sheets carrying 邮编难易度 / 熟手PPH, with their column indexes
DIFFICULTY_SRC = [("LAS01", "LAS线路价格(财务版本）", "LAS", 3, 5, 6),
                  ("TUC01", "TUC线路价格", "TUC", 3, 5, 6)]
GRADE_ORDER = {"A": 0, "B": 1, "C": 2, "D": 3}


def dsp_of(fleet):
    """'PHX-ECS' -> 'ECS';  'XAE' -> 'XAE'."""
    return fleet.split("-", 1)[1] if "-" in fleet else fleet


def parse_grade(raw):
    """'B级 中' -> ('B', '中').  Returns (None, None) when absent."""
    if not raw:
        return None, None
    s = str(raw).strip()
    m = re.match(r"([A-D])\s*级\s*(.*)", s)
    return (m.group(1), m.group(2).strip() or None) if m else (None, s)


def read_difficulty():
    """(station, zip) -> {'grade','grade_label','pph'}"""
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    out = {}
    for station, sheet, srcname, cz, cg, cp in DIFFICULTY_SRC:
        ws = wb[sheet]
        ws.reset_dimensions()
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if r[0] is None or str(r[0]).strip().upper() != srcname:
                continue
            g, lab = parse_grade(r[cg])
            pph = r[cp] if isinstance(r[cp], (int, float)) else None
            out[(station, int(str(r[cz]).strip()))] = {
                "grade": g, "grade_label": lab,
                "pph": round(float(pph), 2) if pph else None}
    wb.close()
    return out


def wavg(pairs):
    """volume-weighted mean of (value, weight); falls back to plain mean."""
    pairs = [(v, w) for v, w in pairs if v and v > 0]
    if not pairs:
        return 0.0
    tw = sum(w for _, w in pairs)
    if tw > 0:
        return sum(v * w for v, w in pairs) / tw
    return sum(v for v, _ in pairs) / len(pairs)


def main():
    aux = json.load(open(os.path.join(HERE, "geo_aux.json"), encoding="utf-8"))
    cent = {int(k): v for k, v in aux["centroids"].items()}
    stcoord = aux["stations"]
    nogeom = set(int(z) for z in aux["missing"])
    diff = read_difficulty()

    wb = openpyxl.load_workbook(GEN, data_only=True)
    addr = {r[0]: (r[1], r[2]) for r in wb["Address"].iter_rows(min_row=2, values_only=True)}
    order = [s for s in wb.sheetnames if s != "Address"]

    stations = {}
    for sn in order:
        rows = list(wb[sn].iter_rows(min_row=2, values_only=True))
        zrecs, dsps = [], set()
        for _, fleet, route, zip_, price, vol in rows:
            fleet = str(fleet).strip()
            code = dsp_of(fleet)
            dsps.add(code)
            d = diff.get((sn, zip_), {})
            lat, lon = cent.get(zip_, (None, None))
            zrecs.append({
                "zip": str(zip_), "volume": vol, "price": round(float(price), 4),
                "dsps": {code: vol}, "routes": {route: vol}, "fleet": fleet,
                "has_geom": zip_ not in nogeom,
                "lat": lat, "lon": lon,
                "coord_source": "zcta" if lat is not None else None,
                "grade": d.get("grade"), "grade_label": d.get("grade_label"),
                "pph": d.get("pph")})

        # PO Box zips have no ZCTA polygon and therefore no centroid. Fall back to
        # the volume-weighted centroid of the station's other zips so that zip
        # search can still fly to them.
        known = [(z["lat"], z["lon"], z["volume"]) for z in zrecs if z["lat"] is not None]
        if known:
            tw = sum(w for _, _, w in known) or len(known)
            wts = [(la, lo, (w or 1)) for la, lo, w in known]
            tw = sum(w for _, _, w in wts)
            flat = sum(la * w for la, _, w in wts) / tw
            flon = sum(lo * w for _, lo, w in wts) / tw
            for z in zrecs:
                if z["lat"] is None:
                    z["lat"], z["lon"] = round(flat, 6), round(flon, 6)
                    z["coord_source"] = "station_service_area"

        # route rollup: volume-weighted price over member zips
        byroute = defaultdict(list)
        for z in zrecs:
            for rt in z["routes"]:
                byroute[rt].append(z)
        routes = []
        for rt, zs in byroute.items():
            dv = Counter()
            for z in zs:
                for d_, v in z["dsps"].items():
                    dv[d_] += v
            routes.append({
                "route": rt,
                "volume": sum(z["volume"] for z in zs),
                "zip_count": len(zs),
                "price": round(wavg([(z["price"], z["volume"]) for z in zs]), 3),
                "dominant_dsp": dv.most_common(1)[0][0] if dv else None})
        routes.sort(key=lambda r: r["route"])

        total = sum(z["volume"] for z in zrecs)
        c = stcoord.get(sn) or {}
        stations[sn] = {
            "name": sn,
            "head": addr[sn][1],
            "address": addr[sn][0],
            "dsps": sorted(dsps),
            "total_volume": total,
            "avg_price": round(wavg([(z["price"], z["volume"]) for z in zrecs]), 3),
            "zip_count": len(zrecs),
            "coord_source": c.get("source", "unknown"),
            "lat": c.get("lat"), "lon": c.get("lon"),
            "zips": sorted(zrecs, key=lambda z: -z["volume"]),
            "routes": routes}
    wb.close()

    # dsp_summary across stations
    dsp_summary = {}
    allz = [(sn, z) for sn, st in stations.items() for z in st["zips"]]
    for code in sorted({d for st in stations.values() for d in st["dsps"]}):
        mine = [(sn, z) for sn, z in allz if code in z["dsps"]]
        dsp_summary[code] = {
            "volume": sum(z["dsps"][code] for _, z in mine),
            "zip_count": len(mine),
            "stations": sorted({sn for sn, _ in mine}),
            "avg_price": round(wavg([(z["price"], z["dsps"][code]) for _, z in mine]), 3)}

    data = {"stations": stations, "dsp_summary": dsp_summary}
    with open(os.path.join(HERE, "mapdata.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, separators=(",", ":"), ensure_ascii=False)

    with open(os.path.join(HERE, "price_adjustments.csv"), "w",
              encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["zip", "price"])
        for _, z in sorted(allz, key=lambda x: x[1]["zip"]):
            w.writerow([z["zip"], f"{z['price']:.2f}"])

    # ---- report ----
    kb = os.path.getsize(os.path.join(HERE, "mapdata.json")) / 1024
    print(f"写出 mapdata.json ({kb:.0f} KB), price_adjustments.csv ({len(allz)} 行)\n")
    print(f"{'站点':<7}{'单量':>7}{'邮编':>5}{'线路':>5}{'均价':>7}  DSP")
    for sn in order:
        s = stations[sn]
        print(f"{sn:<7}{s['total_volume']:>7}{s['zip_count']:>5}{len(s['routes']):>5}"
              f"{s['avg_price']:>7.3f}  {'/'.join(s['dsps'])}")
    tv = sum(s["total_volume"] for s in stations.values())
    print(f"{'合计':<7}{tv:>7}{sum(s['zip_count'] for s in stations.values()):>5}"
          f"{sum(len(s['routes']) for s in stations.values()):>5}"
          f"{wavg([(s['avg_price'], s['total_volume']) for s in stations.values()]):>7.3f}")

    print(f"\nDSP {len(dsp_summary)} 家:")
    for d, v in sorted(dsp_summary.items(), key=lambda kv: -kv[1]["volume"]):
        print(f"   {d:<5}{v['volume']:>7}  {v['zip_count']:>3} 邮编  ${v['avg_price']:.3f}  "
              f"{'/'.join(v['stations'])}")

    g = Counter(z["grade"] for _, z in allz)
    print(f"\n难易度覆盖: {sum(v for k,v in g.items() if k)}/{len(allz)} 个邮编  "
          + ", ".join(f"{k or '无'}级={v}" for k, v in sorted(g.items(), key=lambda kv: GRADE_ORDER.get(kv[0], 9))))
    ng = [z["zip"] for _, z in allz if not z["has_geom"]]
    print(f"无地图边界(仅作标记点): {len(ng)} 个邮编 {ng}")
    fb = [(sn, z["zip"], z["volume"]) for sn, z in allz
          if z["coord_source"] == "station_service_area"]
    print(f"坐标回退为服务区中心: {len(fb)} 个 " + ", ".join(f"{s}/{z}({v}单)" for s, z, v in fb))
    noc = [z["zip"] for _, z in allz if z["lat"] is None]
    print(f"仍无坐标: {len(noc)} 个 {noc or '无'}")

    lats = [z["lat"] for _, z in allz if z["lat"]]
    lons = [z["lon"] for _, z in allz if z["lon"]]
    print(f"\n片区范围: 纬度 {min(lats):.2f}~{max(lats):.2f}, 经度 {min(lons):.2f}~{max(lons):.2f}")
    print(f"建议初始中心: {(min(lats)+max(lats))/2:.2f}, {(min(lons)+max(lons))/2:.2f}")


if __name__ == "__main__":
    main()
