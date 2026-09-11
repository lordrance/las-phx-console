# -*- coding: utf-8 -*-
"""
Fetch 2020 Census ZCTA5 boundaries + centroids for the LAS/PHX region's zips,
and geocode the 8 station addresses. Writes to the scratchpad as JSON.

Boundaries: TIGERweb MapServer layer 2, server-side simplified
(maxAllowableOffset=0.001 deg ~= 100 m, geometryPrecision=5 ~= 1 m).
"""
import json
import os
import time
import urllib.parse
import urllib.request
import warnings

import openpyxl

warnings.filterwarnings("ignore")

BASE = r"C:\Users\uyran\OneDrive\Desktop\rance\nc-ops-console\模板和数据"
GEN = os.path.join(BASE, "拉斯维加斯凤凰城各站点信息.xlsx")
OUTDIR = os.path.dirname(os.path.abspath(__file__))

ZCTA = ("https://tigerweb.geo.census.gov/arcgis/rest/services/TIGERweb/"
        "tigerWMS_Current/MapServer/2/query")
GEOCODE = "https://geocoding.geo.census.gov/geocoder/locations/onelineaddress"
OFFSET = "0.001"
PRECISION = "5"
BATCH = 40


def get(url, timeout=120, tries=3):
    for k in range(tries):
        try:
            return urllib.request.urlopen(url, timeout=timeout).read()
        except Exception as e:
            if k == tries - 1:
                raise
            print(f"      retry {k+1}: {e}")
            time.sleep(3)


def read_table():
    wb = openpyxl.load_workbook(GEN, data_only=True)
    zips, addrs = set(), {}
    for sn in wb.sheetnames:
        if sn == "Address":
            for r in wb[sn].iter_rows(min_row=2, values_only=True):
                addrs[r[0]] = r[1]
            continue
        for r in wb[sn].iter_rows(min_row=2, values_only=True):
            zips.add(str(r[3]))
    wb.close()
    return sorted(zips), addrs


def fetch_boundaries(zips):
    feats, cent = {}, {}
    for i in range(0, len(zips), BATCH):
        chunk = zips[i:i + BATCH]
        p = {"where": "ZCTA5 IN (%s)" % ",".join("'%s'" % z for z in chunk),
             "outFields": "ZCTA5,CENTLAT,CENTLON,INTPTLAT,INTPTLON",
             "returnGeometry": "true", "f": "geojson", "outSR": "4326",
             "maxAllowableOffset": OFFSET, "geometryPrecision": PRECISION}
        raw = get(ZCTA + "?" + urllib.parse.urlencode(p))
        j = json.loads(raw)
        for f in j.get("features", []):
            z = f["properties"]["ZCTA5"]
            feats[z] = {"type": "Feature", "properties": {"zip": z},
                        "geometry": f["geometry"]}
            # INTPTLAT is the internal point: guaranteed inside the polygon
            cent[z] = (float(f["properties"]["INTPTLAT"]),
                       float(f["properties"]["INTPTLON"]))
        print(f"   批次 {i//BATCH+1}: 请求 {len(chunk)} 个, 累计拿到 {len(feats)} 个多边形"
              f" ({len(raw)/1024:.0f} KB)")
    return feats, cent


def geocode(addrs):
    out = {}
    for st, ad in addrs.items():
        # TUC01 holds two sites separated by ';' -- geocode the first (ECS/GOE)
        one = ad.split(";")[0]
        one = one.split("(")[0].strip().rstrip(",")
        p = {"address": one, "benchmark": "Public_AR_Current", "format": "json"}
        try:
            j = json.loads(get(GEOCODE + "?" + urllib.parse.urlencode(p), timeout=90))
            m = j["result"]["addressMatches"]
            if m:
                c = m[0]["coordinates"]
                out[st] = {"lat": c["y"], "lon": c["x"],
                           "matched": m[0]["matchedAddress"], "source": "census"}
                print(f"   OK  {st:<6} {c['y']:.5f}, {c['x']:.5f}  <- {m[0]['matchedAddress']}")
                continue
        except Exception as e:
            print(f"   !!  {st}: {e}")
        out[st] = None
        print(f"   !!  {st:<6} 地理编码失败, 待回退: {one}")
    return out


def main():
    zips, addrs = read_table()
    print(f"表中邮编 {len(zips)} 个, 站点 {len(addrs)} 个\n")

    print("[1] 拉取 ZCTA 邮编边界")
    feats, cent = fetch_boundaries(zips)
    missing = [z for z in zips if z not in feats]
    print(f"\n   拿到 {len(feats)}/{len(zips)} 个多边形")
    if missing:
        print(f"   无边界数据的邮编 ({len(missing)} 个, 多为 PO Box 专用): {missing}")

    print("\n[2] 站点地址地理编码 (Census Geocoder)")
    coords = geocode(addrs)

    fc = {"type": "FeatureCollection", "features": [feats[z] for z in zips if z in feats]}
    with open(os.path.join(OUTDIR, "zipgeo.json"), "w", encoding="utf-8") as f:
        json.dump(fc, f, separators=(",", ":"), ensure_ascii=False)
    with open(os.path.join(OUTDIR, "geo_aux.json"), "w", encoding="utf-8") as f:
        json.dump({"centroids": cent, "stations": coords, "missing": missing},
                  f, indent=1, ensure_ascii=False)

    kb = os.path.getsize(os.path.join(OUTDIR, "zipgeo.json")) / 1024
    print(f"\n写出 zipgeo.json ({kb:.0f} KB) 和 geo_aux.json")
    pts = sum(sum(len(r) for r in g["geometry"]["coordinates"])
              if g["geometry"]["type"] == "Polygon"
              else sum(len(r) for poly in g["geometry"]["coordinates"] for r in poly)
              for g in fc["features"])
    print(f"总坐标点数 {pts:,}, 平均每个邮编 {pts/len(fc['features']):.0f} 点")


if __name__ == "__main__":
    main()
